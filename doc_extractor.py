import mimetypes
from pathlib import Path
import re
import fitz  # PyMuPDF
import docx
import openpyxl
import pandas as pd
import odf.text
import odf.opendocument
from striprtf.striprtf import rtf_to_text
from ocr import OCR


class DocExtractor:
    ocr = OCR()

    @staticmethod
    def extract_paragraphs(file_path):
        file_path = Path(file_path)
        mime_type, _ = mimetypes.guess_type(file_path)

        handler = DocExtractor._get_handlers().get(mime_type)
        if handler:
            return handler(file_path)
        else:
            raise ValueError(f"Unsupported MIME type: {mime_type}")

    @staticmethod
    def extract_text(file_path):
        return '\n\n'.join(paragraph for _, paragraph in DocExtractor.extract_paragraphs(file_path))

    @staticmethod
    def _get_handlers():
        return {
            'application/pdf': DocExtractor._handle_pdf,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': DocExtractor._handle_docx,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': DocExtractor._handle_xlsx,
            'application/vnd.oasis.opendocument.text': DocExtractor._handle_odt,
            'application/rtf': DocExtractor._handle_rtf,
            'text/plain': DocExtractor._handle_plain_text,
            'text/csv': DocExtractor._handle_plain_text,
            'text/xml': DocExtractor._handle_plain_text,
            'application/xml': DocExtractor._handle_plain_text,
            'application/json': DocExtractor._handle_plain_text,
        }

    # === Handlers ===
    
    @staticmethod
    def _handle_pdf(file_path):
        doc = fitz.open(file_path)
        paragraphs = []

        for i, page in enumerate(doc):
            page_num = i + 1

            # Check if the page contains images
            if page.get_images(full=True):
                # Run OCR on the image-containing page
                ocr_page_data = DocExtractor.ocr.extract_pdf_data(file_path, i)
                for block in ocr_page_data.blocks:
                    block_text = OCR._convert_block_data_to_text(block)
                    if block_text.strip():
                        paragraphs.append((page_num, block_text.strip()))
            else:
                # Extract regular text
                text = page.get_text().strip()
                if text:
                    page_paragraphs = DocExtractor._split_into_paragraphs(text)
                    paragraphs.extend((page_num, p) for p in page_paragraphs)

        return paragraphs


    @staticmethod
    def _handle_docx(file_path):
        doc = docx.Document(file_path)
        return [(None, p.text.strip()) for p in doc.paragraphs if p.text.strip()]

    @staticmethod
    def _handle_xlsx(file_path):
        wb = openpyxl.load_workbook(file_path, data_only=True)
        paragraphs = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None]
                if values:
                    paragraphs.append((None, " | ".join(values)))
        return paragraphs

    @staticmethod
    def _handle_odt(file_path):
        textdoc = odf.opendocument.load(str(file_path))
        text_elements = textdoc.getElementsByType(odf.text.P)
        return [(None, str(p).strip()) for p in text_elements if str(p).strip()]

    @staticmethod
    def _handle_rtf(file_path):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
            text = rtf_to_text(file.read())
            return [(None, p) for p in DocExtractor._split_into_paragraphs(text)]

    @staticmethod
    def _handle_plain_text(file_path):
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()
            return [(None, p) for p in DocExtractor._split_into_paragraphs(text)]

    @staticmethod
    def _split_into_paragraphs(text):
        return [p.strip() for p in re.split(r'(?:\r?\n\s*){2,}', text) if p.strip()]
