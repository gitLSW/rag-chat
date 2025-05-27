import re
from pathlib import Path
import mimetypes

from services.ocr import OCR
import fitz  # PyMuPDF
import docx
import openpyxl
import odf.text
import odf.opendocument
from striprtf.striprtf import rtf_to_text


class DocExtractor:
    ocr = OCR()

    @staticmethod
    def extract_paragraphs(file_path, force_ocr=False):
        file_path = Path(file_path)

        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        if not file_path.is_file():
            raise ValueError(f"Not a file: {file_path}")

        mime_type, _ = mimetypes.guess_type(file_path)
        if not mime_type:
            raise ValueError(f"Cannot determine MIME type for file: {file_path}")

        if force_ocr and mime_type == 'application/pdf':
            return DocExtractor._handle_pdf_with_ocr(file_path)

        handler = DocExtractor._get_handlers().get(mime_type)
        if handler:
            try:
                return handler(file_path)                
            except Exception as e:
                raise RuntimeError(f"Failed to extract content from {file_path.name}: {str(e)}") from e
        else:
            raise ValueError(f"Unsupported MIME type: {mime_type}")


    @staticmethod
    def extract_text(file_path):
        return '\n\n'.join(paragraph for _, paragraph in DocExtractor.extract_paragraphs(file_path))


    @staticmethod
    def _get_handlers():
        return {
            'application/pdf': DocExtractor._handle_pdf,
            'application/vnd.openxmlformats-officedocument.wordprocessingml.document': DocExtractor._handle_docx,
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet': DocExtractor._handle_xlsx,
            'application/vnd.oasis.opendocument.text': DocExtractor._handle_odt,
            'application/rtf': DocExtractor._handle_rtf,
            'text/plain': DocExtractor._handle_plain_text,
            'text/csv': DocExtractor._handle_plain_text,
            'text/xml': DocExtractor._handle_plain_text,
            'application/xml': DocExtractor._handle_plain_text,
            'application/json': DocExtractor._handle_plain_text,
        }


    # === Handlers ===

    @staticmethod
    def _handle_pdf(file_path):
        try:
            doc = fitz.open(file_path)
        except Exception as e:
            raise OSError(f"Unable to open PDF: {str(e)}") from e

        paragraphs = []

        for i, page in enumerate(doc):
            page_num = i + 1
            try:
                # Pages with images are read by the OCR
                if page.get_images(full=True):
                    ocr_page_data = DocExtractor.ocr.extract_pdf_data(file_path, i).pages[0]
                    for block in ocr_page_data.blocks:
                        paragraph = OCR._convert_block_data_to_paragraph(block)
                        if paragraph.strip():
                            paragraphs.append((page_num, paragraph.strip()))
                else:
                    text = page.get_text().strip()
                    if text:
                        page_paragraphs = DocExtractor._split_into_paragraphs(text)
                        paragraphs.extend((page_num, p) for p in page_paragraphs)
            except Exception as e:
                raise RuntimeError(f"Failed to process page {page_num} of PDF: {str(e)}") from e

        return paragraphs
    

    @staticmethod
    def _handle_pdf_with_ocr(file_path):
        paragraphs = []
        try:
            doc = DocExtractor.ocr.extract_pdf_data(file_path)
            for i, ocr_page_data in enumerate(doc.pages):
                for block in ocr_page_data.blocks:
                    paragraph = OCR._convert_block_data_to_paragraph(block)
                    if paragraph.strip():
                        paragraphs.append((i, paragraph.strip()))
            return paragraphs
        except Exception as e:
            raise RuntimeError(f"Failed to process PDF with OCR: {str(e)}") from e


    @staticmethod
    def _handle_docx(file_path):
        try:
            doc = docx.Document(file_path)
            return [(None, p.text.strip()) for p in doc.paragraphs if p.text.strip()]
        except Exception as e:
            raise OSError(f"Error reading DOCX: {str(e)}") from e


    @staticmethod
    def _handle_xlsx(file_path):
        try:
            wb = openpyxl.load_workbook(file_path, data_only=True)
            paragraphs = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    values = [str(cell) for cell in row if cell is not None]
                    if values:
                        paragraphs.append((None, ' | '.join(values)))
            return paragraphs
        except Exception as e:
            raise OSError(f"Error reading XLSX: {str(e)}") from e


    @staticmethod
    def _handle_odt(file_path):
        try:
            textdoc = odf.opendocument.load(str(file_path))
            text_elements = textdoc.getElementsByType(odf.text.P)
            return [(None, str(p).strip()) for p in text_elements if str(p).strip()]
        except Exception as e:
            raise OSError(f"Error reading ODT: {str(e)}") from e


    @staticmethod
    def _handle_rtf(file_path):
        try:
            with open(file_path, 'r', errors='ignore') as file:
                text = rtf_to_text(file.read())
                return [(None, p) for p in DocExtractor._split_into_paragraphs(text)]
        except Exception as e:
            raise OSError(f"Error reading RTF: {str(e)}") from e


    @staticmethod
    def _handle_plain_text(file_path):
        try:
            with open(file_path, 'r', errors='ignore') as f:
                text = f.read()
                return [(None, p) for p in DocExtractor._split_into_paragraphs(text)]
        except Exception as e:
            raise OSError(f"Error reading plain text: {str(e)}") from e


    @staticmethod
    def _split_into_paragraphs(text):
        return [p.strip() for p in re.split(r"(?:\r?\n\s*){2,}", text) if p.strip()]